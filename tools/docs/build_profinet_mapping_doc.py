from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(
    r"C:\Users\SIMATIC User\Downloads\SCHLENKER_36_10_ROB_PLC_HMI_IMPLEMENTATION_UPDATE_2026-09-19.docx"
)
OUTPUT = Path(
    r"C:\www\Schlenker\docs\SCHLENKER_36_10_ROB_PROFINET_NETWORK_MAPPING_2026-09-21.docx"
)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def style_run(run, size=8.5, bold=False, color="000000"):
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")


def add_table(doc, headers, rows, widths, font_size=8.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    repeat_header(header)
    for index, text in enumerate(headers):
        cell = header.cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F4E78")
        set_cell_border(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        style_run(paragraph.add_run(text), font_size, True, "FFFFFF")
    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2:
                set_cell_shading(cell, "EAF2F8")
            set_cell_border(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if index in (0, 1, len(values) - 1)
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            style_run(paragraph.add_run(str(value)), font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style="Heading 1" if level == 1 else "Heading 2")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10 if level == 1 else 7)
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run(text)
    style_run(run, 14 if level == 1 else 11.5, True)
    return paragraph


def add_body(doc, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.08
    if bold_lead and text.startswith(bold_lead):
        style_run(paragraph.add_run(bold_lead), 9.5, True)
        style_run(paragraph.add_run(text[len(bold_lead):]), 9.5)
    else:
        style_run(paragraph.add_run(text), 9.5)
    return paragraph


def main():
    if not SOURCE.exists():
        raise FileNotFoundError(SOURCE)
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document(str(SOURCE))

    doc.paragraphs[0].text = (
        "SCHLENKER 36/10\n"
        "PLC HMI IMPLEMENTATION AND PROFINET NETWORK MAPPING FOR ROB"
    )
    doc.paragraphs[1].text = (
        "Revision date: 21 September 2026 | Source requirements: 19 September 2026 | "
        "Project software: TIA Portal V19"
    )
    doc.paragraphs[2].text = (
        "Purpose: coordinate the logical PROFINET, IO-Link and G120C mapping for the current "
        "PLC/HMI project. The software has been compiled offline with zero errors and zero "
        "warnings. This revision does not claim that missing physical field devices, addresses, "
        "telegrams or safety mappings have been commissioned."
    )
    title_p_pr = doc.paragraphs[0]._p.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    add_heading(doc, "11 PROFINET Network Mapping Coordination", 1)
    add_body(
        doc,
        "Decision: the functional allocation is coordinated and may be used for engineering. "
        "Physical PROFINET IP addresses, process-image addresses, GSDML/IODD selections and "
        "G120C telegrams remain release hold points until the installed hardware data is approved.",
        "Decision: ",
    )
    add_body(
        doc,
        "Fail-safe rule: Configured, CommunicationOK and DataValid must remain FALSE for every "
        "uncommissioned node. PLC command bits must never be reused as running or ready feedback.",
        "Fail-safe rule: ",
    )

    add_heading(doc, "11 1 Coordinated Network Nodes", 2)
    network_rows = [
        ("PLC_1", "S7-1512C-1 PN controller", "PLC_1", "192.168.10.1", "PN/IE_1", "VERIFIED PROJECT NODE"),
        ("HMI_1", "MTP1200 Unified Comfort", "HMI_1", "VERIFY FINAL INTERFACE", "X1/X2", "EXISTING PROJECT NODE"),
        ("AL100", "Bottle and accumulation IO-Link", "al100", "TBC", "IFM AL1403", "LOGICAL MAP COMPLETE"),
        ("AL101", "Cap system and air pressure IO-Link", "al101", "TBC", "IFM AL1403", "LOGICAL MAP COMPLETE"),
        ("AL102", "Process instrumentation IO-Link", "al102", "TBC", "IFM AL1403", "DEVICES TBC"),
        ("AL103", "Additional machine sensors", "al103", "TBC", "IFM AL1403", "REQUIRED FLAG TO CONFIRM"),
        ("AL104", "Customer CIP digital interface", "al104", "TBC", "IFM AL1403", "REQUIRED LOGICAL NODE"),
        ("M100", "Main machine drive", "g120c-main", "TBC", "SINAMICS G120C", "TELEGRAM TBC"),
        ("M101", "Bottle conveyor drive", "g120c-conveyor", "TBC", "SINAMICS G120C", "TELEGRAM TBC"),
        ("M102", "Product pump drive", "g120c-product-pump", "TBC", "SINAMICS G120C", "TELEGRAM TBC"),
        ("TBC", "Cap distributor drive", "g120c-cap-distributor", "TBC", "SINAMICS G120C", "CODE AND TELEGRAM TBC"),
        ("SAFE1", "Pilz PNOZmulti diagnostics", "TBC", "TBC", "PROFINET non-safety diagnostics", "PHYSICAL MAP TBC"),
        ("SMC1", "Anybus ABC3113-A to SMC EX260", "TBC", "TBC", "PROFINET to EtherCAT", "PHYSICAL MAP TBC"),
    ]
    add_table(
        doc,
        ("ID", "Function", "Proposed device name", "IP address", "Interface", "Release status"),
        network_rows,
        (0.65, 2.35, 1.75, 1.45, 1.85, 1.65),
        7.8,
    )

    add_heading(doc, "11 2 IO Link Master Coordination", 2)
    master_rows = [
        ("AL100", "Bottle and accumulation sensors", "P1-P4", "P5-P8", "Required", "NOT CONFIGURED"),
        ("AL101", "Cap system and air pressure", "P1-P5", "P6-P8", "Required", "NOT CONFIGURED"),
        ("AL102", "Process instrumentation", "P1-P7 candidates", "P8", "Required", "DEVICES AND SCALING TBC"),
        ("AL103", "Additional machine sensors", "P1-P4 reserved", "P5-P8", "Confirm", "NOT INSTALLED"),
        ("AL104", "Customer CIP interface", "P1-P8", "None", "Required", "PHYSICAL INTERFACE TBC"),
    ]
    add_table(
        doc,
        ("Master", "Function", "Allocated ports", "Spare ports", "Required", "Current status"),
        master_rows,
        (0.8, 2.6, 1.75, 1.25, 1.1, 2.0),
        8.2,
    )

    add_heading(doc, "11 3 IO Link Port Mapping", 2)
    port_rows = [
        ("AL100", "P1", "DI-031", "Bottle Shortage 1", "b_BottleShortage_1", "DI_Bottle_Shortage_1", "FROZEN TBC"),
        ("AL100", "P2", "DI-032", "Bottle Shortage 2", "b_BottleShortage_2", "DI_Bottle_Shortage_2", "FROZEN TBC"),
        ("AL100", "P3", "DI-033", "Outfeed Accumulation 1", "b_Accumulation_1", "DI_Accumulation_1", "FROZEN TBC"),
        ("AL100", "P4", "DI-034", "Outfeed Accumulation 2", "b_Accumulation_2", "DI_Accumulation_2", "FROZEN TBC"),
        ("AL100", "P5-P8", "SPARE", "Future bottle or expansion sensors", "-", "-", "SPARE"),
        ("AL101", "P1", "DI-040", "Cap Present Pick and Place", "b_CapPresent_PP", "DI_Cap_Present", "FROZEN TBC"),
        ("AL101", "P2", "DI-060", "Cap Hopper Low", "b_CapHopper_Low", "DI_Cap_Hopper_Low", "FROZEN TBC"),
        ("AL101", "P3", "DI-061", "Cap Channel Request", "b_CapChannel_Req", "DI_Cap_Channel_Demand", "FROZEN TBC"),
        ("AL101", "P4", "DI-062", "Cap Channel Empty", "b_CapChannel_Empty", "DI_Caps_Missing", "FROZEN TBC"),
        ("AL101", "P5", "DI-010", "Machine Air Pressure OK", "b_AirPressureOK", "DI_AirPressure_OK", "FROZEN TBC"),
        ("AL101", "P6-P8", "SPARE", "Cap pneumatic expansion", "-", "-", "SPARE"),
        ("AL102", "P1", "PT-PRODUCT", "Product Pressure", "ai_ProductPressure", "TBC", "DEVICE TBC"),
        ("AL102", "P2", "PT-AIR", "Air or process pressure", "ai_AirPressure", "TBC", "DEVICE TBC"),
        ("AL102", "P3", "PT-CIP", "CIP service pressure", "ai_CIPPressure", "TBC", "DEVICE TBC"),
        ("AL102", "P4", "PT-SPARE", "Spare pressure", "ai_SparePressure", "TBC", "SPARE"),
        ("AL102", "P5", "LT-PRODUCT", "Product continuous level", "ai_ProductLevel", "TBC", "DEVICE TBC"),
        ("AL102", "P6", "LT-CIP", "CIP process level", "ai_CIPLevel", "TBC", "DEVICE TBC"),
        ("AL102", "P7", "TT-PROCESS", "Process temperature", "ai_ProcessTemperature", "TBC", "DEVICE TBC"),
        ("AL102", "P8", "SPARE", "Future process instrument", "-", "-", "SPARE"),
        ("AL103", "P1-P4", "SEN-103-01..04", "Additional machine sensors", "io_AL103_P1..P4", "TBC", "RESERVED"),
        ("AL103", "P5-P8", "SPARE", "Future smart sensors", "-", "-", "SPARE"),
        ("AL104", "P1", "CIP-DO-01", "Cold Water Request", "CIP_REQ_COLD_WATER", "CIP_Req_Cold_Water", "FROZEN TBC"),
        ("AL104", "P2", "CIP-DO-02", "Hot Water Request", "CIP_REQ_HOT_WATER", "CIP_Req_Hot_Water", "FROZEN TBC"),
        ("AL104", "P3", "CIP-DO-03", "Acid Chemical Request", "CIP_REQ_ACID", "CIP_Req_Acid", "FROZEN TBC"),
        ("AL104", "P4", "CIP-DO-04", "Citra Neutralizer Request", "CIP_REQ_CITRA", "CIP_Req_Citra", "FROZEN TBC"),
        ("AL104", "P5", "CIP-DO-05", "Discharge Request", "CIP_REQ_DISCHARGE", "CIP_Req_Discharge", "FROZEN TBC"),
        ("AL104", "P6", "PROD-DO-01", "Production Product Request", "PRODUCTION_REQ_PRODUCT", "Production_Req_Product", "FROZEN TBC"),
        ("AL104", "P7", "CIP-DI-01", "Customer CIP Ready Accepted", "CIP_REMOTE_READY", "CIP_Remote_Ready", "FEEDBACK TBC"),
        ("AL104", "P8", "CIP-DI-02", "Customer CIP Fault Busy", "CIP_REMOTE_FAULT", "CIP_Remote_Fault", "FEEDBACK TBC"),
    ]
    add_table(
        doc,
        ("Master", "Port", "Device ID", "Function", "PLC logical tag", "HMI tag", "Status"),
        port_rows,
        (0.68, 0.62, 1.25, 2.15, 2.0, 1.8, 1.15),
        7.1,
    )

    add_heading(doc, "11 4 G120C Process Data Coordination", 2)
    drive_rows = [
        ("M100", "Main machine", "DB_Global.Out.MainRun / MainSpeedPct", "MainDriveReady / MainDriveFault / MainSpeedActualPct", "Ready Run Fault Actual speed Reset", "TBC"),
        ("M101", "Bottle conveyor", "DB_Global.Out.ConveyorRun / ConveyorSpeedPct", "ConveyorReady / ConveyorFault / actual speed TBC", "Ready Run Fault Actual speed Reset", "TBC"),
        ("M102", "Product pump", "DB_Global.Out.ProductPumpRun / ProductPumpSpeedPct", "ProductPumpReady / ProductPumpFault / actual speed TBC", "Ready Run Fault Actual speed Reset", "TBC"),
        ("TBC", "Cap distributor vibrator", "DB_Global.Out.CapDriveRun / CapDriveSpeedPct", "CapDriveReady / CapDriveFault / actual speed TBC", "Ready Run Fault Actual speed Reset", "TBC"),
    ]
    add_table(
        doc,
        ("Code", "Drive function", "PLC command reference", "PLC feedback reference", "Required interface", "Telegram"),
        drive_rows,
        (0.65, 1.55, 2.35, 2.45, 2.2, 0.8),
        7.6,
    )
    add_body(
        doc,
        "Each G120C must provide independent communication health, ready, running, warning, "
        "fault, reset acknowledgement and actual-speed validity. A common network-health bit is "
        "not sufficient for individual drive permissives.",
    )

    add_heading(doc, "11 5 Release Hold Points", 2)
    hold_rows = [
        ("IO-Link masters", "Order number and firmware; approved GSDML; device name and IP; input/output lengths"),
        ("IO-Link ports", "Installed sensor model; IODD; mode; byte order; polarity; scaling; quality and diagnostic bits"),
        ("G120C drives", "Order number and firmware; device name and IP; selected telegram; PZD lengths; speed scaling; motor data"),
        ("Pilz", "Approved non-safety diagnostic GSDML and raw byte map; safety validation remains separate"),
        ("Anybus SMC", "Gateway device name and IP; PROFINET and EtherCAT process map; valve channel allocation"),
        ("HMI PLC path", "Confirm final HMI interface and subnet used for PLC communication before field-node addressing"),
    ]
    add_table(doc, ("Area", "Required before physical mapping release"), hold_rows, (2.0, 7.7), 8.3)

    add_heading(doc, "11 6 Controlled Implementation Sequence", 2)
    sequence = [
        "1. Approve the installed device list, device names and collision-free IP plan.",
        "2. Import the exact GSDML and IODD revisions used by the installed hardware.",
        "3. Add AL100-AL104, four G120C nodes, Pilz diagnostics and Anybus SMC to the offline TIA hardware configuration.",
        "4. Assign process-image ranges and record every byte, word, bit, data type and scaling in this mapping.",
        "5. Bind the physical process data to the existing DB_Global and OpenPoints symbolic interfaces.",
        "6. Compile PLC hardware software and HMI offline; retain Configured and DataValid FALSE for incomplete nodes.",
        "7. Request explicit operator confirmation immediately before any download.",
        "8. Perform only 24 VDC communication and diagnostic checks until 400 VAC drive and safety commissioning is complete.",
    ]
    for item in sequence:
        add_body(doc, item)

    add_body(
        doc,
        "Status: LOGICAL MAPPING COORDINATED. PHYSICAL PROFINET ADDRESSING AND TELEGRAM "
        "COMMISSIONING WAITING FOR APPROVED HARDWARE DATA.",
        "Status: ",
    )

    for style_name, size in (("Title", 18), ("Heading 1", 14), ("Heading 2", 11.5)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Arial")
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Arial")
        style_p_pr = style._element.find(qn("w:pPr"))
        style_border = None if style_p_pr is None else style_p_pr.find(qn("w:pBdr"))
        if style_border is not None:
            style_p_pr.remove(style_border)

    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
